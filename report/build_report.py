"""Fill the college report template with the AppScout project report.

Reads the template (and certificates, when given) from private/ (not shared), every number
from results/results.json, and the figures from figures/ and report/assets/.

    python report/build_report.py john       # private/AppScout_DSE_Report_John.docx
    python report/build_report.py leonal     # ... one copy per team member, each with that
    python report/build_report.py jebarson   # member's certificate, cover and marks sheet
"""

from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / "private"
FIG = ROOT / "figures"
ASSETS = ROOT / "report" / "assets"
R = json.loads((ROOT / "results" / "results.json").read_text(encoding="utf-8"))

# (name, register number, certificate image in private/report_assets or None, output file)
MEMBERS = {
    "john": ("JOHN CHRISPIN DURAI A", "URK24CS1060", "cert_john.png", "AppScout_DSE_Report_John.docx"),
    "leonal": ("LEONAL PAUL J", "URK24CS1152", "cert_leonal.png", "AppScout_DSE_Report_Leonal.docx"),
    "jebarson": ("JEBARSON P", "URK24CS1159", "cert_jebarson.png", "AppScout_DSE_Report_Jebarson.docx"),
}
STUDENT, REG_NO, CERTIFICATE, OUT_NAME = MEMBERS[sys.argv[1] if len(sys.argv) > 1 else "john"]
TEAM = [(n, r) for n, r, _, _ in MEMBERS.values()]
HEADING_TITLES = {"INTRODUCTION", "LITERATURE REVIEW", "SYSTEM DESIGN", "IMPLEMENTATION", "TESTING AND VALIDATION",
                  "RESULTS AND DISCUSSION", "ABSTRACT", "CONCLUSION", "REFERENCES", "APPENDICES"}
TITLE = "AppScout: Is Your App Idea Viable?"
SUBTITLE = ("A data science system that judges a Google Play Store app idea against 9,659 apps, "
            "37,427 user reviews and a fresh 2026 scrape of the store")
FONT = "Times New Roman"

D, EX, PR, CL, CU, PV, RC = R["data"], R["explore"], R["predict"], R["classify"], R["cluster"], R["prove"], R["recommend"]
K18, REV, SCR = D["kaggle"], D["reviews"], D["scrape"]
pc = lambda v, d=0: f"{v * 100:.{d}f}%"
f2 = lambda v: f"{v:.2f}"
nice = lambda c: c.replace("_", " ").title().replace(" And ", " & ")
SET18, SET26 = "2018 held-out apps", "2026 apps released after the snapshot"


def fmt_p(p):
    return "< 10⁻¹⁶" if p < 1e-16 else f"{p:.1e}" if p < 1e-3 else f"{p:.3f}"


# ------------------------------------------------------------------ low-level helpers

def set_font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic


def delete(p):
    p._element.getparent().remove(p._element)


def set_text(p, text, size=None, bold=None):
    """Replace a paragraph's visible text, keeping page-break runs and the first run's look."""
    runs = [r for r in p.runs if r.text.strip()]
    if not runs:
        runs = [p.add_run()]
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    if size or bold is not None:
        set_font(runs[0], size or 12, bold=bool(bold))


class Writer:
    """Writes content in order, immediately before an anchor paragraph."""

    def __init__(self, doc, anchor, chapter, bullet_numpr):
        self.doc, self.anchor, self.ch, self.numpr = doc, anchor, chapter, bullet_numpr
        self.fig = self.tab = 0

    def _p(self, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.5, before=0):
        p = self.anchor.insert_paragraph_before()
        p.style = self.doc.styles["Normal"]
        pf = p.paragraph_format
        pf.alignment, pf.space_after, pf.space_before, pf.line_spacing = align, Pt(after), Pt(before), line
        return p

    def h(self, text):
        p = self._p(align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, line=1.15)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(text), 13, bold=True)

    def para(self, text):
        p = self._p()
        for i, part in enumerate(text.split("**")):        # **bold** inline
            if part:
                set_font(p.add_run(part), 12, bold=bool(i % 2))

    def bullets(self, items):
        for text in items:
            p = self._p(after=3, line=1.3)
            p.style = self.doc.styles["List Paragraph"]
            p._p.get_or_add_pPr().insert(0, copy.deepcopy(self.numpr))
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.5)
            for i, part in enumerate(text.split("**")):
                if part:
                    set_font(p.add_run(part), 12, bold=bool(i % 2))

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self._p(align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
            r = p.add_run(line if line else " ")
            r.font.name = "Consolas"; r.font.size = Pt(9.5)

    def figure(self, path, caption, width=15.5):
        self.fig += 1
        p = self._p(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, line=1.0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width))
        c = self._p(align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)
        set_font(c.add_run(f"Figure {self.ch}.{self.fig}: {caption}"), 10.5, italic=True)

    def table(self, header, rows, caption, widths=None):
        self.tab += 1
        c = self._p(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, line=1.0)
        c.paragraph_format.keep_with_next = True
        set_font(c.add_run(f"Table {self.ch}.{self.tab}: {caption}"), 10.5, italic=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080")
            borders.append(e)
        t._tbl.tblPr.append(borders)
        for i, row in enumerate([header] + rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                if widths:
                    cell.width = Cm(widths[j])
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                set_font(p.add_run(str(val)), 10.5, bold=(i == 0))
                if i == 0:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E8EEF4")
                    cell._tc.get_or_add_tcPr().append(shd)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            if i < len(t.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        self.anchor._p.addprevious(t._tbl)
        self._p(after=4, line=1.0)


# ------------------------------------------------------------------ numbers used in several places

ct = CL["category_test"]
st_eng = CL["success_table"]["Logistic: + neighbourhood (engine)"]
sv18, sv26 = PV["score_validation"][SET18], PV["score_validation"][SET26]
cmp_ = PV["score_comparison"]
bands18, bands26 = PV["verdict_bands"][SET18]["success"], PV["verdict_bands"][SET26]["success"]
p_at = RC["precision_at_k"]
best_h, best_l = PR["best_honest"], PR["best_leaked"]
tn = EX["then_now"]
fp = EX["free_paid"]


# ------------------------------------------------------------------ the content

def abstract(w):
    w.para("Every year thousands of people decide to build a mobile app, usually on instinct. Most apps never "
           "reach a meaningful audience: in the Google Play Store data used here, only "
           f"{pc(K18['success_share'])} of apps ever reached a million installs. This project builds **AppScout**, a "
           "data science system that takes an app idea written as one sentence and judges how viable it is from "
           "real Play Store data, explains why, draws the market around it, and suggests better-aligned directions.")
    w.para(f"Three data sources are used: a 2018 scrape of {K18['apps']:,} Play Store apps and {REV['rows_with_text']:,} user "
           f"reviews, and a fresh scrape of {SCR['apps_with_details']:,} apps collected for this project in October 2026 with "
           "the google-play-scraper library. The project follows the full data science workflow of the course: a data "
           "quality audit, exploratory analysis and hypothesis testing, regression, classification, clustering, a "
           "content-based and hybrid recommender, and an ethics review covering transparency, fairness and bias.")
    w.para(f"An idea is placed in its market by a TF-IDF and logistic regression model trained on app names (top-3 "
           f"accuracy {pc(ct['balanced_top3'])} against {pc(ct['majority_top3'])} for always guessing the three largest categories). "
           f"Its 25 most similar existing apps form its neighbourhood, and a logistic regression on what is knowable "
           f"before launch estimates the chance of reaching a million installs (AUC {f2(st_eng['auc'])} on held-out apps). "
           f"The viability score is validated on apps the system never saw: AUC {f2(sv18['auc_score'])} on held-out 2018 apps "
           f"and {f2(sv26['auc_score'])} on apps launched after the 2018 data was collected. Among those new apps, "
           f"{pc(bands26[3])} of ideas rated \"Strong\" reached a million installs against {pc(bands26[0])} of those rated "
           "\"High risk\". A central finding is that the intuitive rule \"few competitors and unhappy users mean "
           "opportunity\" points the wrong way in this data, so the score is learned from data rather than set by hand.")
    w.para("Key features: a web application (Streamlit) where any idea is scored in seconds, with the exact contribution "
           "of each input, the closest competitors from 2018 and live from the Play Store today, the category's install, "
           "rating and pricing charts, better-aligned niches, and the topics that dominate users' complaints; seven "
           "executed analysis notebooks; and one results file from which every number in this report is read.")


def chapter1(w):
    w.h("1.1 Background")
    w.para("Mobile apps are one of the most common products a student or small team can build and publish alone. The "
           "Google Play Store lowers the cost of launching to almost nothing, which means a new app enters a market "
           f"already holding millions of others. In the {K18['apps']:,}-app sample studied here, the median app had "
           f"{10 ** EX['describe']['log_installs']['median']:,.0f} installs, and only about one app in three reached a "
           "million. The decision that matters most, which idea to build, is usually made before any of this is looked at.")
    w.para("The store itself is public. Every listing shows its category, installs, rating, price and reviews, and "
           "researchers have shown that this kind of data can be mined to understand what users want and how apps "
           "succeed [5], [6]. The question this project asks is whether that public data can be turned around and used "
           "**before** an app exists, to judge an idea.")
    w.h("1.2 Problem statement")
    w.para("Given an app idea described in one sentence, with optionally its category, business model, price and "
           "content rating, estimate how viable it is on the Google Play Store, explain the estimate in terms a founder "
           "can act on, and suggest directions that are better aligned with what succeeds.")
    w.h("1.3 Motivation")
    w.bullets(["**Ideas are cheap, building is not.** A few seconds of evidence before months of work is a good trade.",
               "**Market tools describe apps that exist.** Commercial market-intelligence platforms report downloads and "
               "revenue of published apps; none starts from an idea that has no data yet.",
               "**A hard data science problem hides inside.** The strongest predictors of installs (reviews, ratings) only "
               "exist after launch. Building a model that is honest about what can be known beforehand is the core "
               "methodological challenge of this project."])
    w.h("1.4 Objectives")
    w.bullets(["Collect and clean Play Store data, including a fresh 2026 scrape to compare with the 2018 snapshot.",
               "Describe the market: demand, competition, quality, pricing and user sentiment, by category and genre.",
               "Place a free-text idea in its market category, and find its closest existing apps.",
               "Estimate the chance an idea reaches a million installs from pre-launch information only, and validate the "
               "estimate on apps the model never saw, including apps launched years later.",
               "Group the market into segments, and recommend better-aligned niches and the complaints to solve.",
               "Test the main beliefs about the store formally, and check that the system is fair across kinds of apps.",
               "Deliver all of this as a web application anyone can use."])
    w.h("1.5 Scope")
    w.para("The system judges market viability from public store data: demand, competition and the track record of "
           "similar apps. It does not judge the quality of execution, marketing or the team, which Section 4.5 "
           "shows decide most of an app's success. The data are English-language listings from the US store.")
    w.h("1.6 Concepts used")
    w.para("The project follows the course text [4] module by module (Table 1.1).")
    w.table(["Course module", "Concepts", "Where used"],
            [["1 Basics of data science", "Data properties, repositories, tool boxes", "Kaggle data, web scraping, Python stack"],
             ["2 Understanding data", "Import, filter, group, rank, missing values, plots", "Notebook 1, 2; cleaning in common.py"],
             ["3 Statistical inference", "Descriptive statistics, correlation, estimation, hypothesis tests", "Notebook 2, 6"],
             ["4 Machine learning", "Linear, multiple, polynomial, logistic regression, KNN, decision trees, "
              "k-means, agglomerative clustering, model evaluation", "Notebooks 3, 4, 5; the engine"],
             ["5 Recommender systems", "Content-based filtering, hybrid recommenders, evaluation", "Notebook 7; the app"],
             ["6 Ethics", "Bias, fairness, transparency, explainability, accountability [23]", "Notebook 6; the app"]],
            "Course modules and where the project uses them", widths=[3.6, 7.2, 4.8])


def chapter2(w):
    w.h("2.1 App store analysis")
    w.para("Harman, Jia and Zhang [5] proposed treating app stores as a new source for mining software repositories, "
           "showing that the public metadata of thousands of apps (price, rating, downloads, descriptions) reveals "
           "relationships that are invisible in any single app. Martin et al. [6] surveyed the field that followed, "
           "covering hundreds of papers that use app store data for requirements engineering, release planning, "
           "security and testing. Most of this work studies apps that already exist. This project takes the same data "
           "and points it at apps that do not exist yet.")
    w.h("2.2 Mining user reviews")
    w.para("User reviews are a direct record of what users like and dislike. Guzman and Maalej [7] extracted app "
           "features from reviews and attached sentiment to each, so that developers can see which features users "
           "complain about. AppScout uses a simpler, transparent variant suited to a short idea: for each category, "
           "the words that are typical of its negative reviews compared with negative reviews elsewhere (smoothed "
           "log-odds, after Monroe et al. [8]), which separates what users complain about from how they complain.")
    w.h("2.3 Text representation and recommenders")
    w.para("TF-IDF weighting [9] represents a text by its terms, weighted up when frequent in the text and down when "
           "common across all texts; cosine similarity then compares two texts by the angle between their vectors. "
           "Content-based recommenders [10] recommend items whose content resembles what the user is looking at, and "
           "hybrid recommenders [11] combine such similarity with other signals. AppScout's competitor finder is "
           "content-based; its niche suggestions are a hybrid of content similarity and market success.")
    w.h("2.4 Leakage")
    w.para("Kaufman et al. [12] define leakage as information about the target that would not legitimately be "
           "available at prediction time. App data is a textbook case: review counts predict installs almost perfectly "
           "because both measure the same users. A viability model for ideas must exclude every post-launch variable, "
           "and chapter 4 measures how much accuracy that honestly costs.")
    w.h("2.5 Comparison with existing solutions")
    w.table(["Approach", "Starts from", "What it gives", "Limitation for a new idea"],
            [["Commercial market intelligence (Sensor Tower, which acquired data.ai, formerly App Annie, in 2024 [13])",
              "Existing apps", "Download and revenue estimates, rankings", "Paid; needs the app to exist"],
             ["Exploratory notebooks on the Kaggle dataset [1]", "The dataset", "Charts of categories, ratings, prices",
              "Describe the past; no idea input, no validated prediction"],
             ["Searching the store by hand", "Keywords", "A list of competitors", "No measure of demand or success odds"],
             ["AppScout (this project)", "A one-sentence idea", "Validated score, drivers, competitors, market charts, "
              "niches, complaint topics", "2018 training data; execution not measurable"]],
            "AppScout against existing ways of judging an app idea", widths=[4.4, 2.6, 4.6, 4.0])


def chapter3(w):
    w.h("3.1 Architecture")
    w.para("Figure 3.1 shows the system. Three sources are cleaned by one shared module. The viability engine is "
           "fitted on the cleaned data and used in two places: by the analysis notebooks, which measure how well it "
           "works and write every number into results.json, and by the web application, which shows it to users. "
           "Because both use the same engine, the app always gives the numbers the report describes.")
    w.figure(ASSETS / "architecture.png", "System architecture")
    w.h("3.2 Data sources")
    w.table(["Source", "Collected", "Size", "Used for"],
            [["Google Play Store apps [1] (Kaggle, CC BY 3.0)", "Scraped August 2018", f"{K18['raw_rows']:,} rows, {K18['apps']:,} apps",
              "Training and testing every model"],
             ["Google Play Store user reviews [1]", "Scraped August 2018", f"{REV['rows_with_text']:,} reviews, {REV['apps_with_text']} apps",
              "Sentiment, complaint topics"],
             ["2026 scrape (this project) [2]", SCR["scraped_at"][:10], f"{SCR['apps_with_details']:,} apps, {SCR['queries']} searches",
              "Then-and-now comparison; out-of-time test"]],
            "Data sources", widths=[5.2, 3.2, 3.6, 3.6])
    w.h("3.3 Product working flow")
    w.para("Figure 3.2 follows one idea through the engine. The text is cleaned of noise words (\"app\", \"free\", "
           "\"best\"...) and turned into a TF-IDF vector of character 3- to 5-grams, the same representation as every "
           "app name. A logistic regression places it in a category (the user can override it). The 25 apps with the "
           "most similar names become its neighbourhood: their installs, their ratings and how closely they crowd "
           "around the idea are three neighbourhood signals. A second logistic regression turns the category, content "
           "rating, price and the neighbourhood signals into the probability of reaching a million installs. The "
           "viability score is that probability's percentile among all existing apps, judged the same way, and the "
           "verdict bands are its quarters: High risk (below 25), Tough market, Promising, Strong opportunity (75 and above).")
    w.figure(ASSETS / "flow.png", "From idea to verdict")
    w.h("3.4 Design decision: a learned score, not a hand-weighted one")
    w.para("The first design combined four signals with hand-chosen weights: 35% neighbours' installs, 20% low "
           "crowding, 15% low neighbour rating and 30% the learned success probability. This encodes the usual business "
           "intuition that demand, few competitors and unhappy users mean an opportunity. The validation in chapter 5 "
           "showed that two of those signals point the wrong way, so the final engine uses the learned probability alone "
           "and shows the neighbourhood signals as context. The score's breakdown is exact: the logistic regression's "
           "log-odds are split into one term per input, measured against an average app, and the terms add up to the "
           "model's answer.")
    w.h("3.5 Web application design")
    w.para("The page opens with a landing section (Figure 6.1) that states what the tool does and links to the idea form and "
           "to the Colab version. The results are built around one idea: show the user's idea the way the store would list it. The verdict appears "
           "as a store listing, with an icon, the idea as its title, its category, and the row of figures a listing shows "
           "under its name, here the viability score, the chance of a million installs, the installs of similar apps, the "
           "category's success rate and the content rating. Below it, a verdict scale from High risk to Strong opportunity "
           "marks the score. The details sit in four tabs, so a demonstration moves quickly.")
    w.bullets(["**Input:** the idea, category (detected or chosen), free or paid, price, content rating, and whether to "
               "search the live store.",
               "**Verdict:** the listing and the verdict scale described above, with the 35% base rate for comparison.",
               "**Why this score (tab):** a bar per input showing how far it pushes the chance up or down, and the neighbourhood "
               "signals in plain language.",
               "**Make it better (tab):** suggestions generated from the data: niches to angle towards, words "
               "the successful neighbours use, complaint topics to solve, and the free-or-paid evidence.",
               "**Competitors (tab):** the closest apps in the 2018 data and today's live search results.",
               "**The market (tab):** installs reached in the category against the whole store, the rating bar to clear, free "
               "against paid, and every genre placed by competition and demand.",
               "**Data limits:** what each source covers and does not."])


def chapter4(w):
    w.h("4.1 Tools")
    w.table(["Tool", "Purpose"],
            [["Python 3, pandas, NumPy", "Loading, cleaning, grouping and reshaping data"],
             ["scikit-learn [3]", "TF-IDF, regression, classification, KNN, trees, clustering, metrics"],
             ["SciPy", "Hypothesis tests, hierarchical clustering"],
             ["Matplotlib, seaborn", "Every figure in the notebooks and this report"],
             ["google-play-scraper [2]", "The 2026 scrape and the app's live competitor search"],
             ["Streamlit, Altair", "The AppScout web application and its interactive charts"],
             ["Jupyter (nbclient)", "Notebooks executed end to end by one build script"]],
            "Tools", widths=[5.2, 10.4])
    w.h("4.2 Cleaning the 2018 data")
    w.para(f"The raw table has {K18['raw_rows']:,} rows. One row has every value shifted a column to the left (its category "
           f"reads \"1.9\") and was dropped. {K18['duplicate_rows']:,} rows repeat an app already listed: the store was "
           f"scraped more than once, and {K18['apps_in_two_categories']} apps appear under two categories. For each app, "
           f"the row with the most reviews was kept, leaving {K18['apps']:,} apps. Installs, price, size and Android "
           "version are stored as text (\"10,000+\", \"$4.99\", \"19M\", \"4.0.3 and up\") and were parsed into numbers. "
           "Installs is a bucket, not a count, so its lower bound is used, on a log scale.")
    w.para(f"Rating is missing for {K18['missing_rating']:,} apps. These are apps with almost no users: the median rated app "
           f"has {D['median_installs_rated']:,} installs, the median unrated app {D['median_installs_unrated']:,}. The values "
           "are missing for a reason, so they were left missing rather than filled with an average that would invent "
           f"ratings for the apps that have none. {D['price_200_plus']} apps cost $200 or more (up to ${D['max_price']:.0f}); "
           "they are real listings and stay, and every model uses log(1 + price).")
    w.figure(FIG / "see_missing.png", "Missing values in the raw 2018 table", 13)
    w.para(f"The reviews table has {REV['rows']:,} rows, but only {REV['rows_with_text']:,} carry a review, and they cover "
           f"{REV['apps_with_text']} apps. Those apps are much more popular than average (Figure 4.2), so everything learned "
           "from reviews describes established apps: the competition a new idea will meet.")
    w.figure(FIG / "see_reviews_coverage.png", "Reviews cover mostly popular apps", 12)
    w.h("4.3 Scraping the 2026 store")
    w.para(f"For each of the 33 categories, two search phrases of the kind a user types were sent to the store, and every "
           f"hit was looked up for its full listing: {SCR['search_hits']:,} distinct apps, {SCR['failed_lookups']} lookup "
           "failed. Search results are ranked by Google, so the sample favours popular apps "
           f"({pc(SCR['success_share'])} have a million installs, against {pc(K18['success_share'])} in 2018). It is used for "
           "comparison and testing, never for store-wide totals.")
    w.code("""
for cat, phrases in QUERIES.items():                 # e.g. FINANCE: ["budget planner", "stock trading"]
    for q in phrases:
        for h in search(q, n_hits=30, lang="en", country="us"):
            hits.setdefault(h["appId"], (cat, q))
with ThreadPoolExecutor(8) as pool:                   # full listing for every hit
    rows = list(pool.map(details, hits))
""")
    w.h("4.4 Exploring the market")
    top, bot = EX["top_success_categories"], EX["bottom_success_categories"]
    w.para(f"Grouping by category separates competition from demand (Figure 4.3). FAMILY and GAME are the most crowded "
           f"categories; the highest success rates are in {', '.join(nice(c) for c in top[:4])}; the lowest in "
           f"{', '.join(nice(c) for c in bot[-4:])}, which are numerous yet rarely become big.")
    w.figure(FIG / "explore_categories.png", "Competition, demand and quality by category")
    d = EX["describe"]
    w.table(["Variable", "Mean", "Median", "Std. dev.", "Skew"],
            [[n, f"{d[k]['mean']:,.2f}", f"{d[k]['median']:,.2f}", f"{d[k]['std']:,.2f}", f"{d[k]['skew']:.2f}"]
             for n, k in [("Rating", "rating"), ("Installs", "installs"), ("log10 installs", "log_installs"), ("Reviews", "reviews"),
                          ("log10 reviews", "log_reviews"), ("Size (MB)", "size_mb"), ("Price ($)", "price")]],
            "Descriptive statistics of the cleaned 2018 apps", widths=[4, 3, 3, 3, 2.6])
    sp = EX["corr_spearman"]["log_installs"]
    w.para(f"Installs and reviews are so skewed that their means are many times their medians, which is why every model "
           f"works on logarithms. The Spearman correlation of installs with reviews is {sp['log_reviews']:.2f}, almost "
           f"perfect, while with rating it is only {sp['rating']:.2f} (Figure 4.4). Recently updated apps "
           f"(ρ = {sp['days_since_update']:.2f} with days since the last update) and larger apps (ρ = {sp['size_mb']:.2f}) "
           f"have more installs. Paid apps reach far fewer people: {pc(fp['Paid']['success_rate'], 1)} of paid apps reached a "
           f"million installs, against {pc(fp['Free']['success_rate'])} of free ones, with a similar mean rating.")
    w.figure(FIG / "explore_corr.png", "Pearson and Spearman correlation of the numeric variables")
    w.para(f"Between 2018 and 2026, the mean rating of the {tn['categories_compared']} comparable categories rose in "
           f"{tn['rating_up_categories']} of them, and paid apps fell from {pc(tn['paid_share_2018'], 1)} of the store to "
           f"{pc(tn['paid_share_2026'], 1)} of search results: the store moved to free apps that earn from subscriptions and "
           "in-app purchases (Figure 4.5).")
    w.figure(FIG / "explore_2018_2026.png", "Then and now: category ratings and paid share, 2018 against the 2026 sample")
    w.h("4.5 Regression: predicting installs")
    w.para("The target is log10(installs). Data were split 60/20/20 by category, choices (polynomial degree, K, tree depth) "
           "were made on the validation set, and the test set was used once. Two groups of models were compared: models "
           "that use reviews and rating, which exist only after launch, and models that use only what an idea has. "
           "Linear, polynomial and KNN regression were compared with a decision tree [22].")
    rows = [[n, f2(v["rmse"]), f2(v["mae"]), f2(v["r2"])] for n, v in sorted(PR["table"].items(), key=lambda kv: kv[1]["rmse"])]
    w.table(["Model (test set)", "RMSE", "MAE", "R²"], rows, "Every regression model on the same held-out apps (log10 installs)",
            widths=[9.6, 2, 2, 2])
    w.para(f"With reviews, R² reaches {f2(best_l['r2'])}: nearly perfect and useless for an idea. Without them, the best "
           f"honest model ({best_h['model'].lower()}) reaches R² {f2(best_h['r2'])} and cuts the error of the baseline by "
           f"{best_h['rmse_cut_vs_baseline_pct']:.0f}%. Most of what decides installs is invisible before launch; what an idea "
           f"can be judged on is its category and its neighbourhood. The name-based KNN with K = {PR['knn_k']} chosen on "
           "validation (Figure 4.7) is how the engine measures demand.")
    w.figure(FIG / "predict_ranking.png", "Leaked models (grey) against honest pre-launch models")
    w.figure(FIG / "predict_knn_k.png", f"Choosing K for name-based KNN on the validation set: K = {PR['knn_k']}", 10)
    w.h("4.6 Classification")
    w.para(f"**Idea to category.** Logistic regression on TF-IDF of app names, compared on validation with words against "
           f"character n-grams, three regularisation strengths and balanced against plain class weights. Character n-grams "
           f"won. FAMILY, a catch-all holding a fifth of the store, absorbed {pc(ct['unbalanced_family_share'])} of predictions "
           f"without balancing (true share {pc(ct['true_family_share'])}); balanced weights cut this to "
           f"{pc(ct['balanced_family_share'])} for a small loss in accuracy (top-1 {pc(ct['balanced_top1'])} against "
           f"{pc(ct['unbalanced_top1'])}). The balanced model, C = {CL['category_C']:.0f}, is used, with its top three guesses "
           f"shown (top-3 accuracy {pc(ct['balanced_top3'])}). On {ct['apps_2026']:,} app names from 2026 it reached top-1 "
           f"{pc(ct['balanced_top1_2026'])}: the model trained in 2018 still transfers.")
    w.figure(FIG / "classify_category_confusion.png", "Idea-to-category confusion matrix, 10 largest categories", 12)
    rows = [[n, f2(v["auc"]), f2(v["accuracy"]), f2(v["precision"]), f2(v["recall"]), f2(v["f1"])] for n, v in CL["success_table"].items()]
    w.para(f"**Will it reach a million installs?** Four classifiers on pre-launch inputs, K and depth chosen on validation "
           f"by ROC AUC [18] (Table 4.4). Adding the neighbourhood signals raises AUC from "
           f"{f2(CL['success_table']['Logistic: category, content, price']['auc'])} to {f2(st_eng['auc'])}. Logistic "
           "regression is used in the engine: it is the best, and its weights can be read and explained (Figure 4.10).")
    w.table(["Model (test set)", "AUC", "Accuracy", "Precision", "Recall", "F1"], rows,
            "Classifiers for reaching 1M installs, before launch", widths=[6.6, 1.8, 1.9, 1.9, 1.7, 1.7])
    w.figure(FIG / "classify_success.png", "ROC curves and the engine's confusion matrix on held-out apps")
    w.figure(FIG / "classify_weights.png", "The success model's strongest weights (log-odds)", 12)
    w.h("4.7 Clustering market segments")
    w.para(f"Each of the {sum(len(v) for v in CU['members'].values())} genres with at least 20 apps was described by six "
           "standardised numbers (number of apps, success rate, median installs, mean rating, paid share, share of "
           "installs held by its ten biggest apps) and clustered with k-means and agglomerative clustering with Ward linkage [20]. "
           f"Silhouette scores [19] stayed below 0.3 for every k: genres shade into one another rather than forming sharp "
           f"groups. With k capped at 6 so that segments stay few enough to act on, k = {CU['k']} was chosen, and the two "
           f"methods agreed well (adjusted Rand index [21] {CU['ari_kmeans_agglomerative']:.2f}).")
    w.table(["Segment", "Genres", "Members (examples)"],
            [[s, len(m), ", ".join(m[:7]) + ("..." if len(m) > 7 else "")] for s, m in CU["members"].items()],
            "Market segments found by k-means", widths=[4.4, 1.8, 9.4])
    w.figure(FIG / "cluster_opportunity.png", "The opportunity map: every genre by competition and demand, coloured by segment")
    w.h("4.8 The recommender")
    w.para("Similar apps are recommended by cosine similarity of TF-IDF vectors (content-based filtering). Collaborative "
           "filtering is not possible here: the data has no user identities or install histories, only anonymous reviews. "
           "Niche suggestions are a hybrid: the genres where the idea's wider neighbourhood of 100 apps falls, weighted by "
           "each genre's success rate. Complaint topics per category come from the reviews (Table 4.6).")
    comp = RC["complaints"]
    w.table(["Category", "Words most typical of its negative reviews"],
            [[nice(c), ", ".join(comp[c][:6])] for c in ["DATING", "FINANCE", "FOOD_AND_DRINK", "HEALTH_AND_FITNESS", "MEDICAL",
                                                          "SHOPPING", "TRAVEL_AND_LOCAL", "PHOTOGRAPHY"] if c in comp],
            "Complaint topics for selected categories", widths=[4.4, 11.2])
    w.h("4.9 Example real-time application: AppScout")
    w.para("AppScout runs in a browser (streamlit run app/app.py). Figure 4.12 shows it judging \"recipe app that plans "
           "meals from what is in your fridge\". The engine loads once (about 20 seconds), after which every idea is "
           "scored in under a second; the optional live search queries the Play Store for today's top results for the "
           "same words.")
    w.figure(ASSETS / "app_top.png", "AppScout: the idea as a store listing, the verdict scale and why", 14.5)


def chapter5(w):
    w.h("5.1 Testing approach")
    w.bullets(["**One split for everything:** 60% train, 20% validation, 20% test, stratified by category, fixed seed. "
               "Every choice was made on validation; test was used once.",
               "**Out-of-time test:** the 2026 apps released after the 2018 snapshot are ideas the engine could not have "
               "seen, from a different store eight years later.",
               "**Leave-one-out neighbourhoods:** when the engine is trained, each app is described by its neighbours other "
               "than itself, so no app's own installs leak into its features.",
               "**Pre-registered hypothesis tests:** seven hypotheses stated before computing, two-sided, α = 0.05, Holm "
               "correction across all seven [14].",
               "**Self-checks in code:** the cleaning module and the engine each run assertion tests; the web app was run "
               "headlessly with Streamlit's testing framework on several ideas."])
    w.h("5.2 Test cases")
    w.table(["Test", "Expected", "Result"],
            [["Cleaning: shifted row, duplicates, 'Varies with device', '$2.99', '1,000+'", "Dropped / parsed correctly", "Pass"],
             ["Train/val/test split on 20 apps", "12 / 4 / 4", "Pass"],
             ["\"calorie counter and diet tracker\"", "Category Health & Fitness; a 'calorie' app among competitors", "Pass"],
             ["Drivers of any idea", "Baseline + contributions = model log-odds", "Pass (to 0.001)"],
             ["Leave-one-out neighbours", "An app never in its own neighbourhood", "Pass"],
             ["User overrides category", "Engine uses the chosen category", "Pass"],
             ["Web app, default idea and a paid dating idea", "Runs with no exception; all sections render", "Pass"],
             ["Live search with no network", "Section says the store could not be reached", "Handled"]],
            "Test cases", widths=[6.4, 6.4, 2.8])
    w.h("5.3 Hypothesis tests")
    w.para("Two groups were compared with the Mann-Whitney U test [15], several groups with the Kruskal-Wallis test [16], "
           "and confidence intervals were estimated by the bootstrap [17], resampling apps 2,000 times.")
    rows = [[n, fmt_p(t["p_holm"]), "Reject H0" if t["reject_H0"] else "Keep H0", t["effect"]] for n, t in PV["tests"].items()]
    w.table(["Hypothesis", "Holm p", "Decision", "Effect size"], rows, "Seven pre-registered tests (two-sided, Holm-corrected)",
            widths=[5.4, 2.2, 2.4, 5.6])
    b = PV["bootstrap"]
    w.para("All seven null hypotheses are rejected, but with almost 10,000 apps, tiny differences are significant, so the "
           "effect sizes matter more. Paid apps are rated slightly higher (mean difference "
           f"{b['paid_minus_free_mean_rating']['estimate']:.2f} stars, bootstrap 95% CI {b['paid_minus_free_mean_rating']['ci95'][0]:.2f} "
           f"to {b['paid_minus_free_mean_rating']['ci95'][1]:.2f}) yet reach far fewer users. Category explains little of the "
           "variation in rating (ε² in Table 5.2). Recently updated apps reached a million installs "
           f"{b['recent_minus_old_success_rate']['estimate'] * 100:.0f} percentage points more often (CI "
           f"{b['recent_minus_old_success_rate']['ci95'][0] * 100:.0f} to {b['recent_minus_old_success_rate']['ci95'][1] * 100:.0f}). "
           "Unexpectedly, successful apps receive less positive reviews: popular apps attract more critical users.")
    w.h("5.4 Validating the viability score")
    w.para("Existing apps were treated as ideas: each name was typed in as the idea with its category, content rating "
           "and price, and the score was compared with what actually happened. Table 5.3 compares each signal, the old "
           "hand-weighted score and the learned score by AUC, the chance that a successful app outscores an unsuccessful one.")
    w.table(["Signal", "AUC, 2018 held-out", "AUC, 2026 new apps"],
            [[n, f2(v[SET18]), f2(v[SET26])] for n, v in cmp_.items()],
            "Which signals tell winners from losers on unseen apps", widths=[7.6, 4, 4])
    lw = PV["learned_weights"]
    w.para(f"Neighbours' installs carry the signal. Low crowding (AUC {f2(cmp_['low crowding (competition)'][SET18])}) and low "
           f"neighbour ratings ({f2(cmp_['low neighbour rating (room to do better)'][SET18])}) point the wrong way: crowded, "
           "well-rated niches are where apps succeed, because developers crowd into markets that pay. This could be "
           "survivorship in the 2018 data, but the 2026 apps are new entrants and the pattern holds for them too. The "
           f"success model learned it unaided: its standardised weight for crowding is {lw['nb_crowding']:.2f}, against "
           f"{lw['nb_demand']:.2f} for neighbours' installs. The learned score beats the hand-weighted one on both test "
           f"sets ({f2(sv18['auc_score'])} against {f2(cmp_['hand-weighted opportunity score'][SET18])}, and "
           f"{f2(sv26['auc_score'])} against {f2(cmp_['hand-weighted opportunity score'][SET26])}), which is why the engine uses it.")
    w.figure(FIG / "prove_score_comparison.png", "AUC of each signal and both score designs on unseen apps")
    w.para("Each verdict band succeeds more often than the one below it, on both test sets (Figure 5.2): from "
           f"{pc(bands18[0])} to {pc(bands18[3])} on held-out 2018 apps, and from {pc(bands26[0])} to {pc(bands26[3])} on new "
           "2026 apps. The 2026 sample comes from search, which favours successful apps, so its numbers confirm that "
           "the signal transfers but are not an exact accuracy.")
    w.figure(FIG / "prove_score_validation.png", "Share of apps reaching 1M installs by verdict band")
    w.h("5.5 Evaluating the recommender")
    w.para("For each held-out app, the engine recommended its most similar apps from the rest of the store, without "
           "seeing genres; a recommendation is relevant when it shares the app's genre. Precision@5 was "
           f"{pc(p_at['5']['content-based (engine)'])}, against {pc(p_at['5']['popularity'], 1)} for recommending the most popular "
           f"apps and {pc(p_at['5']['random'], 1)} at random; {pc(RC['precision5_same_category'])} of recommendations shared the "
           "category. The niche suggestions cannot be validated the same way, because no app records which niche it should "
           "have chosen; they rest on the demand signal, which is the part of the score that did validate.")
    w.figure(FIG / "recommend_precision.png", "Precision@k of similar-app recommendations against baselines", 12)
    w.h("5.6 Fairness")
    fs = PV["fairness_summary"]
    w.para(f"The success model was checked for each category and for free against paid apps (Figure 5.4). AUC within "
           f"the {fs['groups_with_auc']} groups with enough apps ranged from {f2(fs['auc_min'])} ({nice(fs['auc_min_group'])}) to "
           f"{f2(fs['auc_max'])} ({nice(fs['auc_max_group'])}). The weakest groups are small (tens of test apps), so their "
           f"AUCs are uncertain. The largest calibration gap, {fs['max_abs_calibration_gap'] * 100:.0f} percentage points, is "
           f"also {nice(fs['max_gap_group'])}: the engine is somewhat too gloomy about food apps. The app says so where it "
           "matters by always showing the category's own record beside the score.")
    w.figure(FIG / "prove_fairness.png", "Discrimination (left) and calibration (right) by group")
    w.h("5.7 Validation against the objectives")
    w.table(["Objective", "Evidence"],
            [["Collect and clean data, including a fresh scrape", f"{K18['apps']:,} clean apps; {SCR['apps_with_details']:,} apps scraped in 2026"],
             ["Describe the market", "Chapter 4.4; seven exploration figures"],
             ["Place an idea and find competitors", f"Category top-3 {pc(ct['balanced_top3'])} (baseline {pc(ct['majority_top3'])}); precision@5 {pc(p_at['5']['content-based (engine)'])}"],
             ["Validated pre-launch success estimate", f"AUC {f2(sv18['auc_score'])} (2018), {f2(sv26['auc_score'])} (2026 new apps)"],
             ["Segments and better-aligned niches", f"{CU['k']} segments; hybrid niche recommender in the app"],
             ["Formal tests and fairness", "Seven Holm-corrected tests; per-group AUC and calibration"],
             ["A usable web application", "AppScout (Streamlit), tested headlessly and in a browser"]],
            "Objectives and the evidence that each was met", widths=[6.4, 9.2])


def chapter6(w):
    w.h("6.1 The final system")
    w.para("AppScout opens with a landing page in the style of a consumer app (Figure 6.1): the name set large over the page, "
           "illustrations of what it measures (an app icon, a rating, a million installs, a search), a ticker of the project's "
           "key figures, and two buttons, one to check an idea and one to open the same tool in Google Colab. Below it, three "
           "steps explain the method, each with the measured figure behind it.")
    w.figure(ASSETS / "app_hero.png", "AppScout: the landing page", 15)
    w.para("Figure 4.12 (the top of the results) and Figure 6.2 (the rest) show AppScout after evaluating an idea. The top gives the verdict, the "
           "probability of a million installs against the store's base rate, and how the idea reads to the category "
           "model. The bars under \"Why this score\" are the success model's exact terms for this idea. Four tabs "
           "hold the suggestions and better-aligned niches, the competitors (2018 and live) and four market charts.")
    w.figure(ASSETS / "app_bottom.png", "AppScout: the Make it better and market tabs", 14.5)
    w.h("6.2 Five ideas, end to end")
    cases = RC["cases"]
    w.table(["Idea", "Category", "Score", "Verdict", "P(1M+)", "Best niche"],
            [[i, nice(c["category"]), f"{c['score']:.0f}", c["verdict"], pc(c["p_success"]), c["niches"][0]["genre"] if c["niches"] else "-"]
             for i, c in cases.items()],
            "Case studies", widths=[5.4, 2.8, 1.4, 2.6, 1.6, 2.4])
    w.figure(FIG / "recommend_cases.png", "Why each case-study idea scored what it did")
    pic = cases["AI photo editor that removes backgrounds"]
    dog = cases["dating app for dog owners"]
    w.para(f"The background remover scores highest: its neighbours are photo tools with millions of installs, and the "
           f"successful ones describe themselves with words like {', '.join(pic['winning_words'][:3])}. The dating app for dog "
           f"owners scores lowest: Dating is a crowded, low-success category, and its reviews complain about "
           f"{', '.join(dog['complaints'][:4])}, which points to what a better dating app would have to solve.")
    w.h("6.3 Key findings")
    w.bullets([f"**Leakage is the trap.** Reviews predict installs with R² {f2(best_l['r2'])}; without them, R² is {f2(best_h['r2'])}. "
               "Most of success is invisible before launch, and the app says so.",
               "**Crowded is not bad.** Crowding follows demand; low competition and low ratings more often mark a market "
               "nobody wants. A learned score beat the intuitive hand-weighted one on every test.",
               f"**Launch free.** {pc(fp['Free']['success_rate'])} of free apps reach a million installs, {pc(fp['Paid']['success_rate'], 1)} of "
               f"paid ones; by 2026 paid apps were {pc(tn['paid_share_2026'], 1)} of search results.",
               "**Keep it updated.** Apps updated in the last six months of the snapshot succeeded far more often.",
               "**What a name signals transfers.** Models trained on 2018 still work on 2026 apps."])
    w.h("6.4 Challenges")
    w.bullets(["**Messy data:** a shifted row, duplicate listings, numbers stored as text, ratings missing for a reason.",
               "**A catch-all category:** FAMILY overlaps with games and education; balanced class weights and showing the "
               "top three guesses handled it.",
               "**Short texts:** app names have two or three words; character n-grams matched plurals and variants that whole "
               "words missed.",
               "**A design that failed its test:** the hand-weighted opportunity score was replaced when validation showed two "
               "of its parts point the wrong way.",
               "**Biased samples:** reviews favour popular apps and the 2026 scrape favours search winners; both are stated "
               "wherever they are used."])
    w.h("6.5 Ethics")
    w.para("The ethics review follows the frameworks of the course text on data science ethics [23].")
    w.bullets(["**Transparency and explainability:** no bare number; every score is split into exact contributions and shown "
               "with the real apps behind it.",
               "**Accountability:** every number in the app and this report is written by code into one results file.",
               "**Bias and fairness:** per-category discrimination and calibration were measured and are reported.",
               "**Privacy:** only public listings and anonymous review texts; no user identities are collected or shown. The "
               "2026 scrape was a single run of about 1,700 public listing lookups.",
               "**Honest uncertainty:** the app calls its output a verdict, not a forecast, and states the base rate."])


def conclusion(w):
    w.h("Summary")
    w.para("AppScout turns public Play Store data into a judgement of an app idea that does not exist yet. Using the full "
           "workflow of the course on 2018 and 2026 store data, it places an idea in its market, finds its competitors, "
           "estimates its chance of reaching a million installs, explains that estimate exactly, and suggests better-aligned "
           f"directions. Its score separated successful from unsuccessful apps on unseen data (AUC {f2(sv18['auc_score'])}), "
           f"and still did on apps launched years after its data was collected (AUC {f2(sv26['auc_score'])}).")
    w.h("Achievements")
    w.bullets(["Every course module applied to one real problem, with every number traceable to the code that produced it.",
               "A fresh 2026 scrape of the Play Store, used both to compare markets and to test the models out of time.",
               "An honest treatment of leakage, and a score design changed by evidence rather than intuition.",
               "A working web application with explanations, live competitors and data-driven suggestions."])
    w.h("Limitations")
    w.bullets(["The training data is a 2018 snapshot of apps that survived to be listed; apps that failed and were removed are missing.",
               "A one-sentence idea is matched against app names only; two different ideas with similar names look alike.",
               "Reviews cover 865 mostly popular apps, so complaint topics describe incumbents.",
               "Execution, marketing and timing, which decide most of an app's success, cannot be measured before launch."])
    w.h("Future enhancements")
    w.bullets(["Retrain on a large, regularly refreshed scrape including descriptions, so ideas are matched on what apps do, "
               "not only what they are called.",
               "Scrape current reviews for the idea's live competitors to give up-to-date complaint topics.",
               "Track ideas scored by users and the apps they launch, to evaluate the system prospectively.",
               "Add sentence embeddings as a second text representation and compare them with TF-IDF."])


REFERENCES = [
    "Gupta, L. (2018). Google Play Store Apps [Data set]. Kaggle. https://www.kaggle.com/datasets/lava18/google-play-store-apps. Licensed under CC BY 3.0.",
    "JoMingyu. google-play-scraper: Google Play Store scraper for Python [Software]. MIT License. https://github.com/JoMingyu/google-play-scraper",
    "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    "Igual, L., & Seguí, S. (2017). Introduction to Data Science: A Python Approach to Concepts, Techniques and Applications. Springer.",
    "Harman, M., Jia, Y., & Zhang, Y. (2012). App store mining and analysis: MSR for app stores. In Proc. 9th IEEE Working Conference on Mining Software Repositories (MSR), 108–111.",
    "Martin, W., Sarro, F., Jia, Y., Zhang, Y., & Harman, M. (2017). A survey of app store analysis for software engineering. IEEE Transactions on Software Engineering, 43(9), 817–847. https://doi.org/10.1109/TSE.2016.2630689",
    "Guzman, E., & Maalej, W. (2014). How do users like this feature? A fine grained sentiment analysis of app reviews. In Proc. IEEE 22nd International Requirements Engineering Conference (RE), 153–162.",
    "Monroe, B. L., Colaresi, M. P., & Quinn, K. M. (2008). Fightin' words: Lexical feature selection and evaluation for identifying the content of political conflict. Political Analysis, 16(4), 372–403.",
    "Salton, G., & Buckley, C. (1988). Term-weighting approaches in automatic text retrieval. Information Processing & Management, 24(5), 513–523.",
    "Lops, P., de Gemmis, M., & Semeraro, G. (2011). Content-based recommender systems: State of the art and trends. In Recommender Systems Handbook (pp. 73–105). Springer.",
    "Burke, R. (2002). Hybrid recommender systems: Survey and experiments. User Modeling and User-Adapted Interaction, 12(4), 331–370.",
    "Kaufman, S., Rosset, S., Perlich, C., & Stitelman, O. (2012). Leakage in data mining: Formulation, detection, and avoidance. ACM Transactions on Knowledge Discovery from Data, 6(4), 15:1–15:21.",
    "Sensor Tower (2024, March 18). Sensor Tower acquires market intelligence platform data.ai [Press release]. https://sensortower.com/press/press-release-sensor-tower-acquires-market-intelligence-platform-data-ai",
    "Holm, S. (1979). A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics, 6, 65–70.",
    "Mann, H. B., & Whitney, D. R. (1947). On a test of whether one of two random variables is stochastically larger than the other. Annals of Mathematical Statistics, 18(1), 50–60.",
    "Kruskal, W. H., & Wallis, W. A. (1952). Use of ranks in one-criterion variance analysis. Journal of the American Statistical Association, 47(260), 583–621.",
    "Efron, B. (1979). Bootstrap methods: Another look at the jackknife. The Annals of Statistics, 7, 1–26.",
    "Hanley, J. A., & McNeil, B. J. (1982). The meaning and use of the area under a receiver operating characteristic (ROC) curve. Radiology, 143(1), 29–36.",
    "Rousseeuw, P. J. (1987). Silhouettes: A graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, 20, 53–65.",
    "Ward, J. H. (1963). Hierarchical grouping to optimize an objective function. Journal of the American Statistical Association, 58, 236–244.",
    "Hubert, L., & Arabie, P. (1985). Comparing partitions. Journal of Classification, 2, 193–218.",
    "Breiman, L., Friedman, J. H., Olshen, R. A., & Stone, C. J. (1984). Classification and Regression Trees. Wadsworth.",
    "Martens, D. (2022). Data Science Ethics: Concepts, Techniques and Cautionary Tales. Oxford University Press.",
]


def appendices(w):
    w.h("A. Running the project")
    w.code("""
pip install -r requirements.txt
python export/scrape_2026.py        # optional: refresh the 2026 scrape
python notebooks/build.py           # runs every notebook, rewrites results.json and figures
streamlit run app/app.py            # the AppScout web application
python report/build_report.py       # this report
""")
    w.table(["Notebook", "Contents", "Course modules"],
            [["01_see_data_quality", "Problems in the raw data, cleaning, missing values, outliers, the scrape", "1, 2"],
             ["02_explore_market", "Descriptive statistics, categories, correlation, free vs paid, sentiment, 2018 vs 2026", "2, 3"],
             ["03_predict_installs", "Simple, multiple, polynomial regression, KNN, tree; leakage", "4"],
             ["04_classify", "Idea to category; success: logistic, KNN, tree; ROC; weights", "4, 6"],
             ["05_cluster_markets", "k-means, agglomerative, silhouette, ARI, PCA, opportunity map", "4"],
             ["06_prove_tests_fairness", "Seven Holm-corrected tests, bootstrap, score validation, fairness", "3, 4, 6"],
             ["07_recommend", "Content-based and hybrid recommender, evaluation, complaints, case studies", "5"]],
            "The analysis notebooks", widths=[4.6, 8.6, 2.4])
    w.h("B. Additional figures")
    w.figure(FIG / "explore_distributions.png", "Installs buckets and ratings of 2018 apps")
    w.figure(FIG / "explore_free_paid.png", "Free against paid: installs and ratings")
    w.figure(FIG / "predict_leak.png", "Reviews against installs: the leaked fit", 12)
    w.figure(FIG / "cluster_choose_k.png", "Choosing the number of segments: elbow and silhouette", 13)
    w.figure(FIG / "cluster_dendrogram.png", "Agglomerative clustering of genres (Ward linkage)")
    w.h("C. The core of the engine")
    w.code("""
def evaluate(self, idea, category=None, paid=False, price=0.0, content="Everyone"):
    xq = self.vec.transform([normalise(idea)])        # TF-IDF of the idea
    probs = self.category_model.predict_proba(xq)[0]   # idea -> category
    category = category or classes[probs.argmax()]
    sim, idx = self.neighbours(xq)                     # 25 most similar names
    sig = self.signals(sim, idx)                       # demand, crowding, rating
    row = sig.assign(category=category, content=content,
                     paid=int(paid), log_price=log1p(price))
    p = self.success_model.predict_proba(row)[0, 1]    # P(1M+ installs)
    score = 100 * percentile(p, self.reference["p_success"])
""")


# ------------------------------------------------------------------ assembling the template

GUIDE_TEXT = {
    "A brief summary of the project", "Objectives and scope of the project", "Key features and functionalities",
    "Background information about the project", "Problem statement and motivation",
    "Overview of the concepts used (Micro Processor & concepts in COA)", "Comparison with similar projects or existing solutions",
    "Architecture and high-level design of the system", "Block diagrams (architecture diagram, flow diagram, etc.)",
    "Description of the product working flow with necessary diagram", "Detailed explanation of the implementation process",
    "Example real time application or product", "Description of the testing approach and methodologies used",
    "Test cases and results", "Validation of the system against the requirements", "Presentation of the final system",
    "Evaluation of the project's success in achieving its objectives",
    "Discussion of any challenges faced during the development process", "Summary of the project",
    "Achievements and limitations", "Future enhancements and recommendations",
    "List of all the references and resources used in the project",
}
PLACEHOLDERS = ("<< Font size-12>>", "<<Font- Times new Roman>>")
GUIDE_PREFIXES = ("Review of relevant literature", "IoT module includes")
SECTIONS = [abstract, chapter1, chapter2, chapter3, chapter4, chapter5, chapter6, conclusion, None, appendices]


def main():
    doc = Document(PRIVATE / "Template for Skill based Evaluation(23CS2024).docx")
    paras = doc.paragraphs
    numpr = copy.deepcopy(next(p for p in paras if p.text.strip() == "A brief summary of the project")._p.pPr.numPr)

    # Cover, title page, evaluation sheet
    for p in paras:
        t = p.text.strip()
        if t == "STUDENT NAME (REG.NO)":
            set_text(p, f"{STUDENT} ({REG_NO})")
        elif t == "TITLE":
            set_text(p, TITLE)
        elif t.startswith("<Font Size 18>"):
            set_text(p, SUBTITLE, size=14)
            p.runs[0].italic = True
        elif t.startswith("TEAM MEMBERS NAME"):
            set_text(p, "\n".join(f"{n} ({r})" for n, r in TEAM), size=16, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "Name:":
            p.add_run(STUDENT.title()).font.size = Pt(12)
        elif t == "Register Number:":
            p.add_run(REG_NO).font.size = Pt(12)
        elif t == "Any additional information, diagrams, or code snippets that are relevant but not included in the main report":
            delete(p)
        elif t.startswith("<<The certificate") and CERTIFICATE:
            p.clear()
            p.add_run().add_picture(str(PRIVATE / "report_assets" / CERTIFICATE), width=Cm(16))

    # Chapter headings: one consistent look.
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("CHAPTER ") or t in HEADING_TITLES:
            set_text(p, t, size=14 if t.startswith("CHAPTER") else 16, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = None
            p.paragraph_format.first_line_indent = None
            p._p.get_or_add_pPr()
            if p._p.pPr.numPr is not None:
                p._p.pPr.remove(p._p.pPr.numPr)

    # Content: each "<<write the content here>>" (and the references one) in order.
    anchors = [p for p in doc.paragraphs if p.text.strip() in ("<<write the content here>>", "<<list the reference here with bullet>>")]
    assert len(anchors) == len(SECTIONS), len(anchors)
    for i, (anchor, fn) in enumerate(zip(anchors, SECTIONS)):
        chapter = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6}.get(i, "A" if fn is appendices else "")
        w = Writer(doc, anchor, chapter, numpr)
        if fn is None:                      # references
            for n, ref in enumerate(REFERENCES, 1):
                p = w._p(align=WD_ALIGN_PARAGRAPH.LEFT, after=4, line=1.15)
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.9)
                set_font(p.add_run(f"[{n}]\t{ref}"), 11)
        else:
            fn(w)
        delete(anchor)

    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t in GUIDE_TEXT or t in PLACEHOLDERS or t.startswith(GUIDE_PREFIXES):
            delete(p)

    # One page per chapter: drop the template's scattered page breaks between the abstract and
    # the evaluation sheet, and start every heading on a new page instead.
    body = doc.paragraphs
    start = next(i for i, p in enumerate(body) if p.text.strip() == "ABSTRACT")
    sheet = next(i for i, p in enumerate(body) if p.text.strip() == "School of Computer Science and Technology")
    sheet_top = sheet - 1 if body[sheet - 1]._p.xpath(".//w:drawing") else sheet
    for p in body[start:sheet_top]:
        for br in p._p.xpath('.//w:br[@w:type="page"]'):
            br.getparent().remove(br)
    for p in body[start:sheet_top]:
        t = p.text.strip()
        if t == "ABSTRACT" or t.startswith("CHAPTER ") or t in ("CONCLUSION", "REFERENCES", "APPENDICES"):
            p.paragraph_format.page_break_before = True
        if t.startswith("CHAPTER "):
            p.paragraph_format.keep_with_next = True
    body[sheet_top].paragraph_format.page_break_before = True

    # The template's empty spacer paragraphs leave gaps under headings and can make blank pages.
    for p in body[start:sheet_top]:
        if not p.text.strip() and not p._p.xpath(".//w:drawing") and not p.paragraph_format.page_break_before:
            delete(p)
    for p in doc.paragraphs:
        if p.text.strip() in HEADING_TITLES:
            p.paragraph_format.space_after = Pt(24)

    # Centred lines: the template pads some with spaces and tabs, and its list style indents them.
    for p in doc.paragraphs:
        if p.alignment != WD_ALIGN_PARAGRAPH.CENTER or not p.text.strip():
            continue
        f = p.paragraph_format
        f.left_indent = f.right_indent = f.first_line_indent = Pt(0)
        runs = [r for r in p.runs if not r._r.xpath("./w:br")]
        for r in runs:
            if not r.text.strip():
                r.text = ""
        texts = [r for r in runs if r.text]
        if texts:
            texts[0].text = texts[0].text.lstrip()
            texts[-1].text = texts[-1].text.rstrip()

    out = PRIVATE / OUT_NAME
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
